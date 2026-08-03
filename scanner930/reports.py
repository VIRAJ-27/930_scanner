from __future__ import annotations

import csv
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import STRATEGY_NAME


QUERIES = {
    "OptionTradeBook": """
        SELECT trade_id, trading_date, symbol, entry_tier, option_symbol,
               expiry, strike, lot_size, paper_lots, paper_quantity,
               entry_time, underlying_entry_price, underlying_initial_sl,
               underlying_tp1_target, entry_quote_time, entry_option_ltp,
               entry_option_bid, entry_option_ask, entry_spread_percent,
               tp1_exit_time, tp1_underlying_price, tp1_option_bid,
               tp1_quantity, remaining_quantity, final_exit_time,
               final_underlying_price, final_option_bid, final_quantity,
               final_exit_reason, realized_pnl, option_return_percent,
               status, error
        FROM option_paper_trades
        WHERE trading_date=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY entry_time
    """,
    "TradeBook": """
        SELECT trade_id, trading_date, symbol, setup_number, entry_time,
               entry_price, entry_tier, quantity, initial_sl, tp1_target, tp1_touch_time,
               tp1_exit_time, tp1_exit_price, tp1_quantity,
               final_exit_time, final_exit_price, final_exit_quantity,
               final_exit_reason, realized_pnl, status
        FROM trades
        WHERE trading_date=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY entry_time
    """,
    "SetupLedger": """
        SELECT * FROM setups
        WHERE trading_date=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY event_ts, id
    """,
    "OrderBook": """
        SELECT * FROM orders
        WHERE substr(requested_ts,1,10)=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY requested_ts, id
    """,
    "LiveEvents": """
        SELECT * FROM events
        WHERE substr(event_ts,1,10)=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY event_ts, id
    """,
    "CompletedCandles": """
        SELECT * FROM candles
        WHERE substr(start_ts,1,10)=(SELECT MAX(trading_date) FROM daily_status)
        ORDER BY start_ts, symbol, timeframe
    """,
    "LiveStatus": "SELECT * FROM daily_status ORDER BY trading_date",
    "DailySummary": """
        SELECT trading_date, COUNT(*) AS trades,
               SUM(CASE WHEN realized_pnl > 0 THEN 1 ELSE 0 END) AS winners,
               SUM(CASE WHEN realized_pnl <= 0 THEN 1 ELSE 0 END) AS non_winners,
               ROUND(SUM(COALESCE(realized_pnl, 0)), 2) AS realized_pnl
        FROM trades GROUP BY trading_date ORDER BY trading_date
    """,
    "StockSummary": """
        SELECT symbol, COUNT(*) AS trades,
               SUM(CASE WHEN realized_pnl > 0 THEN 1 ELSE 0 END) AS winners,
               ROUND(SUM(COALESCE(realized_pnl, 0)), 2) AS realized_pnl
        FROM trades GROUP BY symbol ORDER BY symbol
    """,
    "OptionDailySummary": """
        SELECT trading_date,
               SUM(CASE WHEN status!='REJECTED' THEN 1 ELSE 0 END) AS entries,
               SUM(CASE WHEN status='REJECTED' THEN 1 ELSE 0 END) AS rejected,
               SUM(CASE WHEN realized_pnl > 0 THEN 1 ELSE 0 END) AS winners,
               SUM(CASE WHEN realized_pnl < 0 THEN 1 ELSE 0 END) AS losers,
               ROUND(SUM(COALESCE(realized_pnl, 0)), 2) AS realized_pnl
        FROM option_paper_trades
        GROUP BY trading_date ORDER BY trading_date
    """,
    "OptionStockSummary": """
        SELECT symbol,
               SUM(CASE WHEN status!='REJECTED' THEN 1 ELSE 0 END) AS entries,
               SUM(CASE WHEN status='REJECTED' THEN 1 ELSE 0 END) AS rejected,
               SUM(CASE WHEN realized_pnl > 0 THEN 1 ELSE 0 END) AS winners,
               ROUND(SUM(COALESCE(realized_pnl, 0)), 2) AS realized_pnl
        FROM option_paper_trades
        GROUP BY symbol ORDER BY symbol
    """,
}


def export_reports(database_path: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database_path)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for name, query in QUERIES.items():
        cursor = connection.execute(query)
        headers = [item[0] for item in cursor.description]
        rows = cursor.fetchall()
        csv_path = output_dir / f"{name}.csv"
        with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)

        # Completed candles can exceed 100k rows per day. Keep the auditable
        # CSV locally without inflating the emailed dashboard workbook.
        if name == "CompletedCandles":
            continue
        sheet = workbook.create_sheet(name[:31])
        sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        _format_sheet(sheet)

    latest_status = connection.execute(
        """
        SELECT updated_ts FROM daily_status
        ORDER BY updated_ts DESC LIMIT 1
        """
    ).fetchone()
    latest_update_ist = ""
    if latest_status and latest_status[0]:
        latest_update_ist = (
            str(latest_status[0])[:19].replace("T", " ") + " IST"
        )

    dashboard = workbook.create_sheet("Dashboard", 0)
    dashboard["A1"] = "9:30 Normal/Silver Option Paper Scanner"
    dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17365D")
    dashboard.merge_cells("A1:H1")
    dashboard.append([])
    dashboard.append(
        [
            "Option Entries",
            None,
            "Winners",
            None,
            "Option P&L",
            None,
            "Open Positions",
            None,
        ]
    )
    dashboard["B3"] = (
        '=COUNTIFS(\'OptionTradeBook\'!$A$2:$A$100000,"<>",'
        '\'OptionTradeBook\'!$AF$2:$AF$100000,"<>REJECTED")'
    )
    dashboard["D3"] = '=COUNTIF(\'OptionTradeBook\'!$AD$2:$AD$100000,">0")'
    dashboard["F3"] = "=SUM('OptionTradeBook'!$AD$2:$AD$100000)"
    dashboard["H3"] = (
        '=COUNTIF(\'OptionTradeBook\'!$AF$2:$AF$100000,"OPEN")+'
        'COUNTIF(\'OptionTradeBook\'!$AF$2:$AF$100000,"RUNNER_OPEN")'
    )
    dashboard.append([])
    dashboard.append(
        [
            "Last Update (IST)",
            latest_update_ist,
            "Mode",
            (
                "=IFERROR(LOOKUP(2,1/('LiveStatus'!$C$2:$C$1000<>\"\"),"
                "'LiveStatus'!$C$2:$C$1000),\"\")"
            ),
            "WebSocket",
            (
                "=IFERROR(LOOKUP(2,1/('LiveStatus'!$D$2:$D$1000<>\"\"),"
                "'LiveStatus'!$D$2:$D$1000),\"\")"
            ),
            "F&O Stocks",
            (
                "=IFERROR(LOOKUP(2,1/('LiveStatus'!$F$2:$F$1000<>\"\"),"
                "'LiveStatus'!$F$2:$F$1000),0)"
            ),
        ]
    )
    dashboard.append([])
    dashboard.append(["Report", "Purpose"])
    dashboard.append(
        ["OptionTradeBook", "Executed/rejected option paper trades and premium P&L"]
    )
    dashboard.append(
        ["TradeBook", "Normal/Silver G1-high entries, TP1, runners and P&L"]
    )
    dashboard.append(["SetupLedger", "Each Trigger/G1/entry-window outcome"])
    dashboard.append(["OrderBook", "Paper/live broker-order audit trail"])
    dashboard.append(["LiveEvents", "Full chronological strategy event log"])
    dashboard.append(["CompletedCandles", "Local CSV: completed 1m/3m OHLCV and VWAP"])
    dashboard.append(["LiveStatus", "Connection, position and P&L health"])
    dashboard.append(["DailySummary", "Daily trade and P&L totals"])
    dashboard.append(["StockSummary", "Per-stock trade and P&L totals"])
    dashboard.append(["OptionDailySummary", "Daily option entries, rejects and P&L"])
    dashboard.append(["OptionStockSummary", "Per-stock option premium results"])
    for cell in dashboard[3]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True, color="17365D")
    for column in ("B", "D", "F", "H"):
        cell = dashboard[f"{column}3"]
        cell.font = Font(size=14, bold=True, color="17365D")
        cell.alignment = Alignment(horizontal="center")
    dashboard["F3"].number_format = '₹#,##0.00;[Red]-₹#,##0.00'
    for cell in dashboard[5]:
        cell.fill = PatternFill("solid", fgColor="EAF2F8")
    thin = Side(style="thin", color="B4C7E7")
    for row in dashboard.iter_rows(min_row=3, max_row=5, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(bottom=thin)
    _format_sheet(dashboard, filter_table=False)
    dashboard.freeze_panes = "A8"
    dashboard.column_dimensions["A"].width = 18
    dashboard.column_dimensions["B"].width = 34
    for column in ("C", "E", "G"):
        dashboard.column_dimensions[column].width = 16
    for column in ("D", "F", "H"):
        dashboard.column_dimensions[column].width = 18
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_dir / "Scanner930_Dashboard.xlsx")
    connection.close()


def option_daily_message(database_path: Path, trading_day: date) -> str:
    connection = sqlite3.connect(database_path)
    row = connection.execute(
        """
        SELECT
            SUM(CASE WHEN status!='REJECTED' THEN 1 ELSE 0 END),
            SUM(CASE WHEN status='REJECTED' THEN 1 ELSE 0 END),
            SUM(CASE WHEN realized_pnl > 0 THEN 1 ELSE 0 END),
            SUM(CASE WHEN realized_pnl < 0 THEN 1 ELSE 0 END),
            COALESCE(SUM(realized_pnl), 0),
            SUM(CASE WHEN status IN ('OPEN','RUNNER_OPEN') THEN 1 ELSE 0 END)
        FROM option_paper_trades WHERE trading_date=?
        """,
        (trading_day.isoformat(),),
    ).fetchone()
    connection.close()
    entries, rejected, winners, losers, pnl, open_positions = [
        value or 0 for value in row
    ]
    return (
        f"9:30 option paper report — {trading_day.isoformat()}\n"
        f"Strategy: {STRATEGY_NAME}\n"
        f"Entries: {int(entries)} | Rejected: {int(rejected)}\n"
        f"Winners: {int(winners)} | Losers: {int(losers)}\n"
        f"Realized option P&L: ₹{float(pnl):,.2f}\n"
        f"Unresolved option positions: {int(open_positions)}\n"
        "Signals and exits are based on the underlying stock. "
        "Premium fills use ask for entry and bid for exit."
    )


def _format_sheet(sheet, filter_table: bool = True) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    if filter_table:
        sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        header = str(sheet.cell(row=1, column=column).value or "").lower()
        values = [
            str(sheet.cell(row=row, column=column).value or "")
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(11, max(map(len, values), default=0) + 2),
            42,
        )
        if header.endswith("_ts") or header.endswith("_time"):
            sheet.column_dimensions[get_column_letter(column)].width = 27
        if any(key in header for key in ("price", "pnl", "vwap", "_sl", "target")):
            for cell in sheet[get_column_letter(column)][1:]:
                cell.number_format = '₹#,##0.00;[Red]-₹#,##0.00'
        elif "quantity" in header or header in {
            "trades",
            "winners",
            "non_winners",
            "open_positions",
            "subscribed_stocks",
            "volume",
        }:
            for cell in sheet[get_column_letter(column)][1:]:
                cell.number_format = "#,##0"
